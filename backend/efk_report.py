# -*- coding: utf-8 -*-
"""
efk_report.py — παραγωγή της "Αναλυτικής Κατάστασης Παραστατικών Αγοράς
Πετρελαίου Κίνησης και Υπολογισμού του Προς Επιστροφή Ποσού Ε.Φ.Κ."
(ΠΑΡΑΡΤΗΜΑ ΙΙα), σε XLSX και PDF.

Μεταφέρθηκε από το C:\\report-tool (2026-09-26, βλ. CLAUDE.md/DONE.md) — εκεί ζούσε σε
ξεχωριστό πρόγραμμα που άνοιγε τη βάση μέσω SQLite URI mode=ro (cross-process, χωρίς καμία
live σύνδεση με το invoicebook). Εδώ είναι πλέον στην ΙΔΙΑ διεργασία, οπότε παίρνει `db`
(το ήδη εισαγμένο `database` module) και χρησιμοποιεί το κανονικό `db.get_db()` — ίδιο
μοτίβο με το `expvault_export.py` (κανόνας του CLAUDE.md: ειδικό-σκοπού export/report
module, όχι γενικό schema, γι' αυτό ζει εδώ και όχι μέσα στο ίδιο το `database.py`).

Σχεδιασμένο πάνω σε πραγματικό, ήδη κατατεθειμένο παράδειγμα (2019, ΛΑΤΟΜΕΙΑ
ΓΑΛΑΤΙΣΤΑΣ ΑΕ). Δομή: header με στοιχεία επιχείρησης/περιόδου, πίνακας με
Α/Α, ημερομηνία, αριθμό παραστατικού, και δύο ζεύγη στηλών (Αξία +
Ποσότητα σε Λίτρα) — ένα για καύσιμο "που χρησιμοποιήθηκε στην παραγωγή"
και ένα για "σε οχήματα που κινούνται σε οδικούς άξονες". Στο πραγματικό
παράδειγμα όλα τα δεδομένα ήταν στο πρώτο ζεύγος — δεν έχουμε (ακόμα)
διάκριση παραγωγής/οχήματος στο schema μας, οπότε το δεύτερο ζεύγος μένει
πάντα κενό (βλ. session notes / project memory).

ΑΞΙΑ = net_amount (προ ΦΠΑ). ΣΥΝΤΕΛΕΣΤΗΣ ΕΦΚ δίνεται ανά κλήση (κρατικό
ποσοστό, αλλάζει με τον καιρό — ποτέ hardcoded). ΠΟΣΟ ΕΠΙΣΤΡΟΦΗΣ =
ΣΥΝΟΛΙΚΗ ΠΟΣΟΤΗΤΑ (λίτρα) × ΣΥΝΤΕΛΕΣΤΗΣ ÷ 1000 (ο συντελεστής είναι ανά
χιλιόλιτρο) — επιβεβαιωμένο με πραγματικά νούμερα από το δείγμα 2019.

Μόνο γραμμές με `efk_eligible=1` μπαίνουν στην αναφορά — ρητό φίλτρο, ΟΧΙ όλες οι γραμμές
Καυσίμων (βλ. `_fetch_rows`).
"""
import os


def _fetch_rows(db, category, date_from, date_to):
    """ΑΞΙΑ/ποσότητα διαβάζονται από τη ΓΡΑΜΜΗ (tbl_invoice_items), όχι το
    invoice header — ένα τιμολόγιο μπορεί να έχει και μη-καύσιμες γραμμές
    μαζί, οπότε το header's net_amount θα ήταν λάθος (θα έμπαινε στην
    αναφορά ΕΦΚ αξία που δεν αφορά καθόλου καύσιμο)."""
    with db.get_db() as conn:
        rows = conn.execute(
            '''SELECT i.doc_date, i.doc_number, it.value as net_amount, it.quantity as quantity_liters,
                      i.customer_name, i.customer_vat, i.customer_doy, i.customer_address, i.customer_phone
               FROM tbl_invoice_items it
               JOIN tbl_invoices i ON i.id = it.invoice_id
               WHERE it.category = ? AND it.efk_eligible = 1 AND i.doc_date BETWEEN ? AND ?
               ORDER BY i.doc_date''',
            (category, date_from, date_to)
        ).fetchall()
        return [dict(r) for r in rows]


def _format_date_gr(iso_date):
    """'YYYY-MM-DD' -> 'DD/MM/YYYY' (ελληνική καταγραφή). Τιμές που δεν
    ταιριάζουν με το αναμενόμενο σχήμα επιστρέφονται ως έχουν."""
    if not iso_date:
        return iso_date
    parts = iso_date.split('-')
    if len(parts) != 3:
        return iso_date
    y, m, d = parts
    return f'{d}/{m}/{y}'


def _safe_filename_part(s):
    return ''.join(c for c in s if c not in '<>:"/\\|?*').strip()


def _build_filenames(output_dir, date_from, date_to):
    label = _safe_filename_part(f'ΕΦΚ {date_from} - {date_to}')
    return (
        os.path.join(output_dir, f'{label}.xlsx'),
        os.path.join(output_dir, f'{label}.pdf'),
    )


def _write_xlsx(path, rows, company, date_from, date_to, rate, total_liters, refund):
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = 'ΕΦΚ'

    bold = Font(bold=True)
    center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    thin = Side(style='thin')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws['A1'] = 'ΠΑΡΑΡΤΗΜΑ ΙΙα'
    ws['A1'].font = bold
    ws['A2'] = f'Για την χρονική περίοδο από {_format_date_gr(date_from)} έως {_format_date_gr(date_to)}'
    ws['A3'] = 'ΑΝΑΛΥΤΙΚΗ ΚΑΤΑΣΤΑΣΗ ΠΑΡΑΣΤΑΤΙΚΩΝ ΑΓΟΡΑΣ ΠΕΤΡΕΛΑΙΟΥ ΚΙΝΗΣΗΣ ΚΑΙ ΥΠΟΛΟΓΙΣΜΟΥ ΤΟΥ ΠΡΟΣ ΕΠΙΣΤΡΟΦΗ ΠΟΣΟΥ Ε.Φ.Κ.'

    ws['A5'] = 'Επωνυμία:'
    ws['B5'] = company.get('name', '')
    ws['A6'] = 'Διεύθυνση:'
    ws['B6'] = company.get('address', '')
    ws['A7'] = 'Τηλέφωνο:'
    ws['B7'] = company.get('phone', '')
    ws['A8'] = 'ΑΦΜ:'
    ws['B8'] = company.get('vat', '')
    ws['A9'] = 'ΔΟΥ:'
    ws['B9'] = company.get('doy', '')
    for r in range(5, 10):
        ws[f'A{r}'].font = bold

    header_row = 11
    headers = [
        'Α/Α', 'ΗΜΕΡΟΜΗΝΙΑ', 'ΑΡΙΘΜΟΣ ΠΑΡΑΣΤΑΤΙΚΟΥ',
        'ΑΞΙΑ (παραγωγή)', 'ΠΟΣΟΤΗΤΑ σε ΛΙΤΡΑ (παραγωγή)',
        'ΑΞΙΑ (οδικοί άξονες)', 'ΠΟΣΟΤΗΤΑ σε ΛΙΤΡΑ (οδικοί άξονες)',
    ]
    for col, h in enumerate(headers, start=1):
        c = ws.cell(row=header_row, column=col, value=h)
        c.font = bold
        c.alignment = center
        c.border = border

    r = header_row + 1
    for i, row in enumerate(rows, start=1):
        values = [
            i, _format_date_gr(row.get('doc_date')), row.get('doc_number'),
            row.get('net_amount'), row.get('quantity_liters'),
            None, None,
        ]
        for col, v in enumerate(values, start=1):
            c = ws.cell(row=r, column=col, value=v)
            c.border = border
        r += 1

    r += 1
    ws.cell(row=r, column=1, value='ΣΥΝΟΛΙΚΗ ΠΟΣΟΤΗΤΑ').font = bold
    ws.cell(row=r, column=5, value=total_liters)
    r += 1
    ws.cell(row=r, column=1, value='ΣΥΝΤΕΛΕΣΤΗΣ ΠΡΟΣ ΕΠΙΣΤΡΟΦΗ Ε.Φ.Κ.').font = bold
    ws.cell(row=r, column=5, value=f'{rate} ευρώ/χιλιόλιτρο')
    r += 1
    ws.cell(row=r, column=1, value='ΠΟΣΟ ΕΠΙΣΤΡΟΦΗΣ').font = bold
    ws.cell(row=r, column=5, value=round(refund, 2))

    widths = [6, 14, 16, 14, 20, 16, 20]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb.save(path)


def _register_greek_font():
    """Οι ενσωματωμένες γραμματοσειρές (Helvetica κλπ) του ReportLab
    χρησιμοποιούν WinAnsiEncoding, που ΔΕΝ καλύπτει τονισμένα ελληνικά
    (ή, ά, ό, κλπ. βγαίνουν ως ■). Καταχωρούμε Arial (TrueType, πλήρες
    Unicode/Greek coverage) — ήδη διαθέσιμο στο Windows."""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    fonts_dir = os.path.join(os.environ.get('WINDIR', r'C:\Windows'), 'Fonts')
    pdfmetrics.registerFont(TTFont('Arial', os.path.join(fonts_dir, 'arial.ttf')))
    pdfmetrics.registerFont(TTFont('Arial-Bold', os.path.join(fonts_dir, 'arialbd.ttf')))
    pdfmetrics.registerFontFamily('Arial', normal='Arial', bold='Arial-Bold')


def _write_pdf(path, rows, company, date_from, date_to, rate, total_liters, refund):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet

    _register_greek_font()
    styles = getSampleStyleSheet()
    styles['Title'].fontName = 'Arial-Bold'
    styles['Normal'].fontName = 'Arial'
    os.makedirs(os.path.dirname(path), exist_ok=True)
    doc = SimpleDocTemplate(path, pagesize=landscape(A4),
                             leftMargin=1.5 * cm, rightMargin=1.5 * cm,
                             topMargin=1.5 * cm, bottomMargin=1.5 * cm)

    elements = [
        Paragraph('ΠΑΡΑΡΤΗΜΑ ΙΙα', styles['Title']),
        Paragraph(f'Για την χρονική περίοδο από {_format_date_gr(date_from)} έως {_format_date_gr(date_to)}', styles['Normal']),
        Paragraph('ΑΝΑΛΥΤΙΚΗ ΚΑΤΑΣΤΑΣΗ ΠΑΡΑΣΤΑΤΙΚΩΝ ΑΓΟΡΑΣ ΠΕΤΡΕΛΑΙΟΥ ΚΙΝΗΣΗΣ ΚΑΙ ΥΠΟΛΟΓΙΣΜΟΥ ΤΟΥ ΠΡΟΣ ΕΠΙΣΤΡΟΦΗ ΠΟΣΟΥ Ε.Φ.Κ.', styles['Normal']),
        Spacer(1, 0.4 * cm),
        Paragraph(f"Επωνυμία: {company.get('name','')} &nbsp;&nbsp; ΑΦΜ: {company.get('vat','')} &nbsp;&nbsp; ΔΟΥ: {company.get('doy','')}", styles['Normal']),
        Paragraph(f"Διεύθυνση: {company.get('address','')} &nbsp;&nbsp; Τηλέφωνο: {company.get('phone','')}", styles['Normal']),
        Spacer(1, 0.5 * cm),
    ]

    data = [['Α/Α', 'ΗΜΕΡΟΜΗΝΙΑ', 'ΑΡ. ΠΑΡΑΣΤΑΤΙΚΟΥ',
             'ΑΞΙΑ\n(παραγωγή)', 'ΠΟΣΟΤΗΤΑ ΛΙΤΡΑ\n(παραγωγή)',
             'ΑΞΙΑ\n(οδικοί άξονες)', 'ΠΟΣΟΤΗΤΑ ΛΙΤΡΑ\n(οδικοί άξονες)']]
    for i, row in enumerate(rows, start=1):
        data.append([
            str(i), _format_date_gr(row.get('doc_date')) or '', str(row.get('doc_number') or ''),
            f"{row.get('net_amount'):.2f}" if row.get('net_amount') is not None else '',
            f"{row.get('quantity_liters'):.2f}" if row.get('quantity_liters') is not None else '',
            '', '',
        ])

    table = Table(data, repeatRows=1)
    table.setStyle(TableStyle([
        ('FONTNAME', (0, 0), (-1, -1), 'Arial'),
        ('FONTNAME', (0, 0), (-1, 0), 'Arial-Bold'),
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0f2040')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
    ]))
    elements.append(table)
    elements.append(Spacer(1, 0.5 * cm))
    elements.append(Paragraph(f'ΣΥΝΟΛΙΚΗ ΠΟΣΟΤΗΤΑ: {total_liters:.2f} λίτρα', styles['Normal']))
    elements.append(Paragraph(f'ΣΥΝΤΕΛΕΣΤΗΣ ΠΡΟΣ ΕΠΙΣΤΡΟΦΗ Ε.Φ.Κ.: {rate} ευρώ/χιλιόλιτρο', styles['Normal']))
    elements.append(Paragraph(f'<b>ΠΟΣΟ ΕΠΙΣΤΡΟΦΗΣ: {refund:.2f} €</b>', styles['Normal']))

    doc.build(elements)


def _single_value(rows, field, label):
    """Επιστρέφει την ΜΙΑ τιμή ενός πεδίου δικαιούχου (π.χ. ΔΟΥ) αν είναι
    συνεπής σε όλες τις εγγραφές — κενές τιμές αγνοούνται (π.χ. τηλέφωνο
    συχνά λείπει από το χαρτί, δεν είναι σύγκρουση). Αν βρεθούν 2+
    ΔΙΑΦΟΡΕΤΙΚΕΣ μη-κενές τιμές, δεν μαντεύουμε ποια είναι σωστή — σφάλμα."""
    values = {(r.get(field) or '').strip() for r in rows}
    values.discard('')
    if len(values) > 1:
        raise ValueError(
            f'Η επιλεγμένη περίοδος έχει ασύμφωνες τιμές για {label}: '
            f'{", ".join(sorted(values))}. Έλεγξε/διόρθωσε τις εγγραφές στο invoicebook.'
        )
    return next(iter(values)) if values else ''


def _derive_beneficiary(rows):
    """Ο 'Δικαιούχος' της αναφοράς ΕΦΚ (Επωνυμία/ΑΦΜ/ΔΟΥ/Διεύθυνση/Τηλέφωνο)
    ΕΙΝΑΙ ο πελάτης των τιμολογίων — ήδη μέσα στα δεδομένα κάθε τιμολογίου
    (τα ίδια πεδία γράφει και το χαρτί), όχι κάτι που πληκτρολογεί ξεχωριστά
    ο χειριστής. Το Επωνυμία/ΑΦΜ ζεύγος είναι η ταυτότητα — αν η περίοδος
    περιέχει τιμολόγια από ΠΑΝΩ ΑΠΟ ΕΝΑΝ πελάτη (π.χ. λόγω αλλαγής νομικού
    προσώπου στο μεσοδιάστημα — έχει ξανασυμβεί, βλ. project memory), δεν
    μαντεύουμε ποιον να διαλέξουμε — σφάλμα, ο χειριστής πρέπει να στενέψει
    την περίοδο."""
    entities = {(r.get('customer_name') or '', r.get('customer_vat') or '') for r in rows}
    entities.discard(('', ''))
    if len(entities) > 1:
        names = ', '.join(f'{n} (ΑΦΜ {v})' for n, v in sorted(entities))
        raise ValueError(
            f'Η επιλεγμένη περίοδος περιέχει τιμολόγια από περισσότερους από έναν πελάτες/'
            f'δικαιούχους: {names}. Στένεψε την περίοδο ώστε να αφορά έναν δικαιούχο.'
        )
    name, vat = next(iter(entities)) if entities else ('', '')
    return {
        'name': name,
        'vat': vat,
        'doy': _single_value(rows, 'customer_doy', 'ΔΟΥ'),
        'address': _single_value(rows, 'customer_address', 'Διεύθυνση'),
        'phone': _single_value(rows, 'customer_phone', 'Τηλέφωνο'),
    }


FUEL_CATEGORY = 'Καύσιμα'


def generate(db, date_from, date_to, rate_per_kiloliter, company, output_dir):
    rows = _fetch_rows(db, FUEL_CATEGORY, date_from, date_to)
    beneficiary = _derive_beneficiary(rows)
    company = dict(company)
    company['name'] = beneficiary['name']
    company['vat'] = beneficiary['vat']
    company['doy'] = beneficiary['doy']
    company['address'] = beneficiary['address']
    company['phone'] = beneficiary['phone']

    total_liters = sum(r.get('quantity_liters') or 0 for r in rows)
    refund = total_liters * float(rate_per_kiloliter) / 1000.0

    xlsx_path, pdf_path = _build_filenames(output_dir, date_from, date_to)
    _write_xlsx(xlsx_path, rows, company, date_from, date_to, rate_per_kiloliter, total_liters, refund)
    _write_pdf(pdf_path, rows, company, date_from, date_to, rate_per_kiloliter, total_liters, refund)

    return {
        'xlsx_path': xlsx_path,
        'pdf_path': pdf_path,
        'row_count': len(rows),
        'total_liters': total_liters,
        'refund_amount': round(refund, 2),
        'beneficiary_name': beneficiary['name'],
        'beneficiary_vat': beneficiary['vat'],
        'beneficiary_doy': beneficiary['doy'],
        'beneficiary_address': beneficiary['address'],
        'beneficiary_phone': beneficiary['phone'],
    }
